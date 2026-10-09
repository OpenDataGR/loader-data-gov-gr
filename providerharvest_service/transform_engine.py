"""The Transform Box's engine: turns a received file into clean records with a list of steps.

Nothing here knows about any one source. A transform is ``reader step -> more steps -> output``:

* a **reader** step understands the *format* of what was received (XML, JSON, CSV) and produces flat
  records -- ``list[dict]``. Which format a source sends is configuration, found from a sample;
* every other **step** works on those flat records and does not care where they came from: add columns
  from an uploaded index table (``lookup``), ``filter``, ``rename``, ``convert`` types, ``select`` or
  ``drop`` fields, ``add_field``;
* the **output** writes the records as JSON, JSON Lines, XML or CSV.

So a source with a different structure needs a different *configuration*, not different code, and a new
step or format is one more entry in ``READERS`` / ``STEPS``. Field paths are the same everywhere:
``a/b/c`` (XML local names -- namespaces and prefixes ignored; JSON keys), with ``@attr`` for an XML
attribute and, for CSV, the column name.

This runs on files that came from outside, so it is bounded and defensive: XML with a DTD or entities is
refused (``defusedxml``), input size, record counts, step counts and index-table sizes are capped, steps are
data (never code: no eval, no templates beyond ``{field}`` substitution), and any problem is a
``TransformError`` with a message fit to show an admin -- never a crash.
"""

from __future__ import annotations

import csv
import dataclasses
import datetime
import io
import json
import re
from collections import Counter
from typing import Any, Callable
from xml.etree import ElementTree as ET

from defusedxml import ElementTree as DET
from defusedxml.common import DefusedXmlException

MAX_INPUT_BYTES = 100 * 1000 * 1000
MAX_RECORDS = 500_000
MAX_STEPS = 20
MAX_FIELDS = 200
MAX_LOOKUP_KEYS_SHOWN = 20
PREVIEW_RECORDS = 50
_NAME = re.compile(r"^[^\W\d]\w{0,63}$")  # letters (any script, Greek included), digits and underscores; not starting with a digit
_PATH = re.compile(r"^(?:[^/@\s]+(?:/[^/@\s]+)*)?(?:@[A-Za-z_][\w.-]*)?$")

Record = dict[str, Any]


class TransformError(ValueError):
    """Carries a message fit to show the person configuring or reading the transform."""


# -- helpers --------------------------------------------------------------------------------------------

def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _text(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text if text != "" else None


def _check_name(name: object, what: str) -> str:
    if not isinstance(name, str) or not _NAME.match(name):
        raise TransformError("%s must be letters, digits and underscores, starting with a letter -- no spaces or symbols (got %r)." % (what, name))
    return name


def _check_path(path: object, what: str) -> str:
    if not isinstance(path, str) or not path or not _PATH.match(path):
        raise TransformError("%s is not a valid path (%r). Use names separated by '/', with @attribute at the end for XML." % (what, path))
    return path


def _fields(options: dict, key: str = "fields", *, required: bool = True, plain: bool = False) -> list[tuple[str, str]]:
    raw = options.get(key) or []
    if not isinstance(raw, list) or (required and not raw):
        raise TransformError("Choose at least one field.")
    if len(raw) > MAX_FIELDS:
        raise TransformError("At most %d fields." % MAX_FIELDS)
    out, seen = [], set()
    for item in raw:
        if not isinstance(item, dict):
            raise TransformError("A field needs a name and a path.")
        name = _check_name(item.get("name"), "A field name")
        path = item.get("path")
        if plain:  # a CSV column: any text
            if not isinstance(path, str) or not path.strip() or len(path) > 200 or any(ord(ch) < 32 for ch in path):
                raise TransformError("Choose the column for %r." % name)
        else:
            path = _check_path(path, "The path of %r" % item.get("name"))
        if name in seen:
            raise TransformError("The field name %r is used twice." % name)
        seen.add(name)
        out.append((name, path))
    return out


# -- XML ----------------------------------------------------------------------------------------------------

def _parse_xml(content: bytes) -> ET.Element:
    try:
        return DET.fromstring(content.lstrip(b"\xef\xbb\xbf"), forbid_dtd=True)
    except DefusedXmlException:
        raise TransformError("This XML uses a DTD or entities, which are refused for safety.") from None
    except ET.ParseError as exc:
        raise TransformError("The file is not valid XML: %s" % exc) from None


def _xml_find(element: ET.Element, path: str) -> list[ET.Element | str]:
    """Everything ``path`` reaches under ``element``: elements, or (for a trailing @attr) attribute values."""
    attr = None
    if "@" in path:
        path, attr = path.rsplit("@", 1)
        path = path.rstrip("/")
    nodes = [element]
    for segment in [s for s in path.split("/") if s]:
        nodes = [child for node in nodes for child in node if _local(child.tag) == segment]
        if not nodes:
            return []
    if attr is None:
        return nodes
    return [value for node in nodes for key, value in node.attrib.items() if _local(key) == attr]


def _xml_value(element: ET.Element, path: str) -> str | None:
    found = _xml_find(element, path)
    if not found:
        return None
    first = found[0]
    return _text(first if isinstance(first, str) else first.text)


def read_xml(content: bytes, options: dict) -> list[Record]:
    record_name = str(options.get("record") or "").strip()
    if not record_name or "/" in record_name or "@" in record_name:
        raise TransformError("Choose the repeating element that holds one record (its name, without any path).")
    fields = _fields(options)
    explode = options.get("explode") or None
    explode_fields = _fields(explode, required=False) if explode else []
    if explode:
        _check_path(explode.get("path"), "The path to the repeated element")
    root = _parse_xml(content)
    records: list[Record] = []
    for element in root.iter():
        if _local(element.tag) != record_name:
            continue
        base = {name: _xml_value(element, path) for name, path in fields}
        if not explode:
            records.append(base)
        else:
            children = [c for c in _xml_find(element, explode["path"]) if not isinstance(c, str)]
            if not children and explode.get("keep_empty"):
                records.append({**base, **{name: None for name, _ in explode_fields}})
            for child in children:
                records.append({**base, **{name: _xml_value(child, path) for name, path in explode_fields}})
        if len(records) > MAX_RECORDS:
            raise TransformError("More than %d records: refusing to continue." % MAX_RECORDS)
    return records


# -- JSON -----------------------------------------------------------------------------------------------------

def _json_load(content: bytes) -> Any:
    try:
        return json.loads(content.decode("utf-8-sig"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise TransformError("The file is not valid JSON: %s" % exc) from None


def _json_find(value: Any, path: str) -> list[Any]:
    nodes = [value]
    for segment in [s for s in path.split("/") if s]:
        nodes = [n[segment] for n in nodes if isinstance(n, dict) and segment in n]
        if not nodes:
            return []
    return nodes


def _json_scalar(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def read_json(content: bytes, options: dict) -> list[Record]:
    data = _json_load(content)
    record_path = str(options.get("record") or "")
    found = _json_find(data, record_path) if record_path else [data]
    if not found:
        raise TransformError("No %r in this JSON." % record_path)
    items = found[0]
    if isinstance(items, dict):
        items = [items]
    if not isinstance(items, list):
        raise TransformError("The records path does not lead to a list of objects.")
    fields = _fields(options, required=False)
    explode = options.get("explode") or None
    explode_fields = _fields(explode, required=False) if explode else []
    records: list[Record] = []
    for item in items:
        if not isinstance(item, dict):
            continue
        if fields:
            base = {name: (lambda f: _json_scalar(f[0]) if f else None)(_json_find(item, path)) for name, path in fields}
        else:
            base = {k: _json_scalar(v) for k, v in item.items() if _NAME.match(str(k))}
        if not explode:
            records.append(base)
        else:
            children = (_json_find(item, explode.get("path", "")) or [[]])[0]
            children = [c for c in children if isinstance(c, dict)] if isinstance(children, list) else []
            if not children and explode.get("keep_empty"):
                records.append({**base, **{name: None for name, _ in explode_fields}})
            for child in children:
                extra = {name: (lambda f: _json_scalar(f[0]) if f else None)(_json_find(child, path)) for name, path in explode_fields}
                records.append({**base, **extra})
        if len(records) > MAX_RECORDS:
            raise TransformError("More than %d records: refusing to continue." % MAX_RECORDS)
    return records


# -- CSV -------------------------------------------------------------------------------------------------------

def read_csv(content: bytes, options: dict) -> list[Record]:
    encoding = str(options.get("encoding") or "utf-8-sig")
    delimiter = str(options.get("delimiter") or ",")
    if len(delimiter) != 1:
        raise TransformError("The delimiter must be one character.")
    try:
        text = content.decode(encoding)
    except (LookupError, UnicodeDecodeError) as exc:
        raise TransformError("Cannot read the file as %s: %s" % (encoding, exc)) from None
    reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
    if not reader.fieldnames:
        raise TransformError("The CSV has no header row.")
    fields = _fields(options, required=False, plain=True)
    columns = fields or [(re.sub(r"\W", "_", c).strip("_") or "col", c) for c in reader.fieldnames]
    records: list[Record] = []
    for row in reader:
        records.append({name: _text(row.get(column)) for name, column in columns})
        if len(records) > MAX_RECORDS:
            raise TransformError("More than %d records: refusing to continue." % MAX_RECORDS)
    return records


# -- Excel (.xlsx) --------------------------------------------------------------------------------------------------------------

MAX_XLSX_UNCOMPRESSED = 200 * 1000 * 1000  # an .xlsx is a zip: refuse one that expands to something absurd
MAX_XLSX_ENTRIES = 500
_OLE_MAGIC = b"\xd0\xcf\x11\xe0"


def _is_xlsx(content: bytes, filename: str = "") -> bool:
    name = (filename or "").lower()
    return name.endswith((".xlsx", ".xlsm")) or content[:4] == b"PK\x03\x04"


def _cell_text(value: Any) -> str | None:
    """A spreadsheet cell as text: whole numbers without a ".0", dates as ISO, blanks as None."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float):
        value = int(value) if value.is_integer() and abs(value) < 1e15 else float(format(value, ".15g"))
    if isinstance(value, datetime.datetime):
        # Excel has no date-only type: a date arrives as midnight, and reads better as just the date.
        return value.date().isoformat() if value.time() == datetime.time(0, 0) and value.tzinfo is None else value.isoformat()
    if isinstance(value, (datetime.date, datetime.time)):
        return value.isoformat()
    return _text(value)


def _open_xlsx(content: bytes):
    import zipfile

    if content[:4] == _OLE_MAGIC:
        raise TransformError("This looks like an old .xls file. Save it as .xlsx (or CSV) and upload that.")
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_ENTRIES or sum(i.file_size for i in infos) > MAX_XLSX_UNCOMPRESSED:
                raise TransformError("This Excel file is too large when unpacked.")
    except zipfile.BadZipFile:
        raise TransformError("This is not a valid .xlsx file (an old .xls, or a damaged file?). Save it as .xlsx or CSV.") from None
    try:
        import openpyxl

        return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except TransformError:
        raise
    except Exception:  # noqa: BLE001 -- openpyxl raises many kinds for damaged files
        raise TransformError("This Excel file could not be read. Is it damaged or password-protected? Try saving it as CSV.") from None


def xlsx_sheets(content: bytes) -> list[str]:
    workbook = _open_xlsx(content)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


def read_xlsx_table(content: bytes, sheet: str | None = None, *, max_rows: int = MAX_RECORDS + 1) -> tuple[str, list[list[str | None]]]:
    """(sheet name used, rows of text) from one sheet: the named one, or the first that has anything in it."""
    workbook = _open_xlsx(content)
    try:
        names = list(workbook.sheetnames)
        if sheet:
            if sheet not in names:
                raise TransformError("There is no sheet %r in this file (it has: %s)." % (sheet, ", ".join(names)))
            candidates_ = [sheet]
        else:
            candidates_ = names
        fallback = None
        for name in candidates_:
            rows: list[list[str | None]] = []
            for row in workbook[name].iter_rows(values_only=True):
                cells = [_cell_text(c) for c in row]
                if any(c is not None for c in cells):
                    rows.append(cells)
                if len(rows) > max_rows:
                    break
            if len(rows) >= 2 or (sheet and rows):
                return name, rows  # a header and something under it (or the sheet the person named)
            if rows and fallback is None:
                fallback = (name, rows)  # a lone row: used only if nothing better turns up
        if fallback:
            return fallback
        raise TransformError("The sheet is empty." if sheet else "This Excel file has no data in any sheet.")
    finally:
        workbook.close()


def _header_and_rows(table: list[list[str | None]]) -> tuple[list[str], list[list[str | None]]]:
    """The first row as column names (a column with no name and no data is dropped; data without a name is an error)."""
    header, body = [(c or "").strip() for c in table[0]], table[1:]
    keep = []
    for index, name in enumerate(header):
        has_data = any(index < len(r) and r[index] is not None for r in body)
        if name:
            keep.append(index)
        elif has_data:
            raise TransformError("Column %d has values but no name in the header row." % (index + 1))
    return [header[i] for i in keep], [[(r[i] if i < len(r) else None) for i in keep] for r in body]


def read_xlsx(content: bytes, options: dict) -> list[Record]:
    _, table = read_xlsx_table(content, str(options.get("sheet") or "") or None)
    columns, rows = _header_and_rows(table)
    fields = _fields(options, required=False, plain=True)
    picked = fields or [(re.sub(r"\W", "_", c).strip("_") or "col", c) for c in columns]
    index = {c: i for i, c in enumerate(columns)}
    records: list[Record] = []
    for row in rows:
        records.append({name: (row[index[column]] if column in index else None) for name, column in picked})
        if len(records) > MAX_RECORDS:
            raise TransformError("More than %d records: refusing to continue." % MAX_RECORDS)
    return records


READERS: dict[str, dict[str, Any]] = {
    "xml_records": {"label": "XML", "read": read_xml},
    "json_records": {"label": "JSON", "read": read_json},
    "csv_records": {"label": "CSV", "read": read_csv},
    "xlsx_records": {"label": "Excel", "read": lambda content, options: read_xlsx(content, options)},
}


def detect_format(content: bytes, filename: str = "") -> str | None:
    """Which reader a file suggests: by its name, then by what it starts with."""
    name = (filename or "").lower()
    for suffix, reader in ((".xml", "xml_records"), (".json", "json_records"), (".csv", "csv_records"), (".xlsx", "xlsx_records"), (".xlsm", "xlsx_records")):
        if name.endswith(suffix):
            return reader
    if content[:4] in (b"PK\x03\x04", b"\xd0\xcf\x11\xe0"):
        return "xlsx_records"
    head = content.lstrip(b"\xef\xbb\xbf \t\r\n")[:1]
    if head == b"<":
        return "xml_records"
    if head in (b"{", b"["):
        return "json_records"
    return None


# -- looking at a sample (so a person can choose records and fields instead of typing paths) ----------------------------------

def candidates(content: bytes, reader: str) -> list[dict]:
    """The places in a sample that look like "one record, repeated": [{"path", "count"}], biggest first."""
    if reader == "xml_records":
        counts = Counter(_local(e.tag) for e in _parse_xml(content).iter())
        return [{"path": n, "count": c} for n, c in counts.most_common() if c > 1][:20]
    if reader == "json_records":
        found: list[dict] = []

        def walk(value: Any, path: str, depth: int) -> None:
            head = value[:5] if isinstance(value, list) else []
            if head and sum(isinstance(v, dict) for v in head) * 2 >= len(head):
                found.append({"path": path, "count": len(value)})
            if depth < 5:
                if isinstance(value, dict):
                    for k, v in value.items():
                        walk(v, (path + "/" if path else "") + str(k), depth + 1)

        data = _json_load(content)
        walk(data, "", 0)
        return sorted(found, key=lambda c: -c["count"])[:20]
    return []


def sniff_delimiter(content: bytes) -> str:
    """The delimiter a CSV's header line most likely uses."""
    first = (content[:4096].decode("utf-8-sig", "replace").splitlines() or [""])[0]
    best = max(",;	|", key=first.count)
    return best if first.count(best) else ","


def describe_record(content: bytes, reader: str, record: str, sample: int = 50, delimiter: str | None = None, sheet: str | None = None) -> dict:
    """The fields available inside a record: {"fields": [{"path", "example", "repeats"}], "repeating": [path...]}.
    ``repeats`` marks a path that occurs more than once in a single record (a candidate to "explode")."""
    seen: dict[str, dict] = {}
    repeating: Counter = Counter()

    def note(path: str, value: Any, per_record: Counter) -> None:
        per_record[path] += 1
        entry = seen.setdefault(path, {"path": path, "example": None, "repeats": False})
        if entry["example"] is None and _text(value) is not None:
            entry["example"] = str(_text(value))[:60]

    if reader == "xml_records":
        root = _parse_xml(content)
        taken = 0
        for element in root.iter():
            if _local(element.tag) != record:
                continue
            taken += 1
            per: Counter = Counter()

            def walk(node: ET.Element, prefix: str) -> None:
                for key, value in node.attrib.items():
                    note(prefix + "@" + _local(key), value, per)
                for child in node:
                    path = (prefix + "/" if prefix else "") + _local(child.tag)
                    if len(child) == 0 and not child.attrib and _text(child.text) is not None:
                        note(path, child.text, per)
                    elif len(child) == 0 and child.attrib:
                        for key, value in child.attrib.items():
                            note(path + "@" + _local(key), value, per)
                        if _text(child.text) is not None:
                            note(path, child.text, per)
                    walk(child, path)

            walk(element, "")
            for path, n in per.items():
                if n > 1:
                    seen[path]["repeats"] = True
                    repeating[path.split("@")[0]] += 1
            if taken >= sample:
                break
    elif reader == "json_records":
        data = _json_load(content)
        items = (_json_find(data, record)[0] if record else data) if (not record or _json_find(data, record)) else []
        for item in (items if isinstance(items, list) else [items])[:sample]:
            per = Counter()

            def walk_json(node: Any, prefix: str) -> None:
                if isinstance(node, dict):
                    for key, value in node.items():
                        walk_json(value, (prefix + "/" if prefix else "") + str(key))
                elif isinstance(node, list):
                    per_path = prefix
                    for value in node[:3]:
                        if isinstance(value, dict):
                            for key, inner in value.items():
                                note(per_path + "/" + str(key), inner, per)
                    if len(node) > 1 and any(isinstance(v, dict) for v in node):
                        repeating[prefix] += 1
                else:
                    note(prefix, node, per)

            if isinstance(item, dict):
                walk_json(item, "")
    elif reader == "csv_records":
        rows = csv.reader(io.StringIO(content.decode("utf-8-sig", "replace"), newline=""), delimiter=delimiter or sniff_delimiter(content))
        header = [c.strip() for c in next(rows, [])]
        first_rows = [r for _, r in zip(range(sample), rows)]
        for index, column in enumerate(header):
            if column:
                seen[column] = {"path": column, "repeats": False, "example": next((r[index][:60] for r in first_rows if index < len(r) and r[index].strip()), None)}
    elif reader == "xlsx_records":
        _, table = read_xlsx_table(content, sheet, max_rows=sample + 1)
        columns, rows = _header_and_rows(table)
        for index, column in enumerate(columns):
            seen[column] = {"path": column, "repeats": False, "example": next((str(r[index])[:60] for r in rows if r[index] is not None), None)}
    else:
        raise TransformError("Unknown format %r." % reader)
    return {"fields": list(seen.values())[:MAX_FIELDS], "repeating": [p for p, _ in repeating.most_common(10)]}


# -- the steps after the reader ------------------------------------------------------------------------------------------------

@dataclasses.dataclass
class Context:
    lookups: dict[str, dict]
    stats: dict[str, Any] = dataclasses.field(default_factory=lambda: {
        "lookup_missing": 0, "missing_keys": [], "dropped_by_lookup": 0, "dropped_by_filter": 0, "conversion_errors": {},
    })


def _number(value: Any) -> float | None:
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _convert(value: Any, kind: str) -> Any:
    if value is None:
        return None
    if kind == "string":
        return str(value)
    if kind == "int":
        number = _number(value)
        if number is None or number != int(number):
            raise ValueError
        return int(number)
    if kind == "float":
        number = _number(value)
        if number is None:
            raise ValueError
        return number
    if kind == "bool":
        text = str(value).strip().lower()
        if text in ("1", "true", "yes", "y", "t"):
            return True
        if text in ("0", "false", "no", "n", "f"):
            return False
        raise ValueError
    if kind == "datetime":
        text = str(value).strip().replace("Z", "+00:00").replace("z", "+00:00")
        moment = datetime.datetime.fromisoformat(text)
        if moment.tzinfo is None:
            moment = moment.replace(tzinfo=datetime.timezone.utc)
        return moment.astimezone(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
    raise ValueError


CONVERT_TYPES = ("string", "int", "float", "bool", "datetime")
FILTER_OPS = ("eq", "ne", "in", "not_in", "gt", "ge", "lt", "le", "contains", "exists", "not_exists")


def step_lookup(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    table_name = options.get("table")
    table = ctx.lookups.get(table_name)
    if table is None:
        raise TransformError("The index table %r does not exist (upload it in the Transform Box)." % table_name)
    key = _check_name(options.get("key"), "The field to match on")
    table_key = str(options.get("table_key") or "")
    add = options.get("add") or []
    if table_key not in table["columns"]:
        raise TransformError("The index table %r has no column %r." % (table_name, table_key))
    if not isinstance(add, list) or not add:
        raise TransformError("Choose at least one column to add from the index table.")
    pairs = []
    for item in add:
        column = str(item.get("column") or "")
        if column not in table["columns"]:
            raise TransformError("The index table %r has no column %r." % (table_name, column))
        pairs.append((column, _check_name(item.get("as") or column, "The new field name")))
    on_missing = options.get("on_missing", "keep")
    if on_missing not in ("keep", "drop"):
        raise TransformError("On a missing match: keep or drop.")
    index = table["index"].get(table_key)
    if index is None:
        index = {}
        for row in table["rows"]:
            value = _text(row.get(table_key))
            if value is not None:
                index.setdefault(value, row)
        table["index"][table_key] = index
    out: list[Record] = []
    missing_seen: set = set()
    for record in records:
        value = _text(record.get(key))
        row = index.get(value) if value is not None else None
        if row is None:
            ctx.stats["lookup_missing"] += 1
            if value is not None and value not in missing_seen and len(ctx.stats["missing_keys"]) < MAX_LOOKUP_KEYS_SHOWN:
                missing_seen.add(value)
                ctx.stats["missing_keys"].append(value)
            if on_missing == "drop":
                ctx.stats["dropped_by_lookup"] += 1
                continue
            out.append({**record, **{new: None for _, new in pairs}})
        else:
            out.append({**record, **{new: _text(row.get(column)) for column, new in pairs}})
    return out


def step_filter(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    field = _check_name(options.get("field"), "The field to test")
    op = options.get("op")
    if op not in FILTER_OPS:
        raise TransformError("Unknown test %r." % op)
    value = options.get("value")
    values = [v.strip() for v in str(value).split(",")] if op in ("in", "not_in") else None
    if op not in ("exists", "not_exists") and (value is None or str(value) == ""):
        raise TransformError("Give a value to compare with.")

    def keep(record: Record) -> bool:
        actual = record.get(field)
        if op == "exists":
            return actual is not None
        if op == "not_exists":
            return actual is None
        if actual is None:
            return op in ("ne", "not_in")
        text = str(actual)
        if op == "eq":
            return text == str(value)
        if op == "ne":
            return text != str(value)
        if op == "in":
            return text in values
        if op == "not_in":
            return text not in values
        if op == "contains":
            return str(value) in text
        a, b = _number(actual), _number(value)
        if a is not None and b is not None:
            return {"gt": a > b, "ge": a >= b, "lt": a < b, "le": a <= b}[op]
        return {"gt": text > str(value), "ge": text >= str(value), "lt": text < str(value), "le": text <= str(value)}[op]

    kept = [r for r in records if keep(r)]
    ctx.stats["dropped_by_filter"] += len(records) - len(kept)
    return kept


def step_rename(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    mapping = options.get("map")
    if not isinstance(mapping, dict) or not mapping:
        raise TransformError("Give at least one field to rename.")
    pairs = {_check_name(old, "A field name"): _check_name(new, "The new name") for old, new in mapping.items()}
    out = []
    for record in records:
        renamed = {pairs.get(k, k): v for k, v in record.items()}
        if len(renamed) != len(record):
            raise TransformError("Renaming makes two fields share a name.")
        out.append(renamed)
    return out


def step_convert(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    spec = options.get("fields")
    if not isinstance(spec, dict) or not spec:
        raise TransformError("Choose at least one field to convert.")
    for field, kind in spec.items():
        _check_name(field, "A field name")
        if kind not in CONVERT_TYPES:
            raise TransformError("Unknown type %r (use %s)." % (kind, ", ".join(CONVERT_TYPES)))
    out = []
    for record in records:
        converted = dict(record)
        for field, kind in spec.items():
            if field in converted:
                try:
                    converted[field] = _convert(converted[field], kind)
                except (ValueError, OverflowError):
                    converted[field] = None
                    ctx.stats["conversion_errors"][field] = ctx.stats["conversion_errors"].get(field, 0) + 1
        out.append(converted)
    return out


def step_select(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    names = [_check_name(n, "A field name") for n in (options.get("fields") or [])]
    if not names:
        raise TransformError("Choose the fields to keep.")
    return [{n: r.get(n) for n in names} for r in records]


def step_drop(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    names = {_check_name(n, "A field name") for n in (options.get("fields") or [])}
    if not names:
        raise TransformError("Choose the fields to drop.")
    return [{k: v for k, v in r.items() if k not in names} for r in records]


_PLACEHOLDER = re.compile(r"\{([A-Za-z_][A-Za-z0-9_]*)\}")


def step_add_field(records: list[Record], options: dict, ctx: Context) -> list[Record]:
    name = _check_name(options.get("name"), "The new field name")
    kind = options.get("kind", "constant")
    value = options.get("value")
    if kind not in ("constant", "template") or value is None or str(value) == "":
        raise TransformError("Give a value (constant text, or a template such as {id}-{time}).")
    out = []
    for record in records:
        if kind == "constant":
            added = str(value)
        else:
            added = _PLACEHOLDER.sub(lambda m: "" if record.get(m.group(1)) is None else str(record.get(m.group(1))), str(value))
        out.append({**record, name: added})
    return out


STEPS: dict[str, dict[str, Any]] = {
    "lookup": {"label": "Add columns from an index table", "apply": step_lookup},
    "filter": {"label": "Keep only matching records", "apply": step_filter},
    "rename": {"label": "Rename fields", "apply": step_rename},
    "convert": {"label": "Convert field types", "apply": step_convert},
    "select": {"label": "Keep only these fields", "apply": step_select},
    "drop": {"label": "Remove fields", "apply": step_drop},
    "add_field": {"label": "Add a field", "apply": step_add_field},
}


def validate(steps: object) -> list[str]:
    """Problems with a list of steps that can be seen without data (empty list: fine)."""
    problems: list[str] = []
    if not isinstance(steps, list) or not steps:
        return ["A transform needs at least a reader step."]
    if len(steps) > MAX_STEPS:
        return ["At most %d steps." % MAX_STEPS]
    for number, step in enumerate(steps, 1):
        kind = step.get("type") if isinstance(step, dict) else None
        if number == 1 and kind not in READERS:
            problems.append("Step 1 must read the file (XML, JSON or CSV).")
        elif number > 1 and kind not in STEPS:
            problems.append("Step %d: unknown step %r." % (number, kind))
        elif number == 1:
            try:
                if kind == "xml_records" and not str(step.get("record") or "").strip():
                    raise TransformError("Choose the repeating element that holds one record.")
                _fields(step, required=kind == "xml_records", plain=kind in ("csv_records", "xlsx_records"))
                if step.get("explode"):
                    _check_path(step["explode"].get("path"), "The path to the repeated element")
                    _fields(step["explode"], required=False)
            except (TransformError, AttributeError) as exc:
                problems.append("Step 1: %s" % exc)
    return problems


# -- output ----------------------------------------------------------------------------------------------------------------------------

OUTPUTS = {
    "json": ("application/json", "json", "JSON (a list of records)"),
    "jsonl": ("application/x-ndjson", "jsonl", "JSON Lines (one record per line)"),
    "xml": ("application/xml", "xml", "XML"),
    "csv": ("text/csv", "csv", "CSV"),
}


def write_output(records: list[Record], fmt: str, *, root: str = "records", record: str = "record") -> bytes:
    if fmt == "json":
        return json.dumps(records, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    if fmt == "jsonl":
        return "\n".join(json.dumps(r, ensure_ascii=False, separators=(",", ":")) for r in records).encode("utf-8")
    if fmt == "csv":
        names: list[str] = []
        for r in records:
            for k in r:
                if k not in names:
                    names.append(k)
        buffer = io.StringIO()
        writer = csv.writer(buffer, lineterminator="\n")
        writer.writerow(names)
        for r in records:
            writer.writerow(["" if r.get(n) is None else r.get(n) for n in names])
        return buffer.getvalue().encode("utf-8")
    if fmt == "xml":
        top = ET.Element(_check_name(root, "The XML root name"))
        for r in records:
            item = ET.SubElement(top, _check_name(record, "The XML record name"))
            for k, v in r.items():
                child = ET.SubElement(item, _check_name(k, "A field name"))
                if v is None:
                    child.set("nil", "true")
                else:
                    child.text = str(v)
        return ET.tostring(top, encoding="utf-8", xml_declaration=True)
    raise TransformError("Unknown output format %r." % fmt)


# -- running a whole transform -------------------------------------------------------------------------------------------------------------------

@dataclasses.dataclass
class Result:
    content: bytes
    content_type: str
    extension: str
    records_in: int
    records_out: int
    stats: dict
    preview: list[Record]


def run(content: bytes, steps: list[dict], *, output: str = "json", lookups: dict[str, dict] | None = None,
        root: str = "records", record: str = "record") -> Result:
    """Apply ``steps`` to a received file. Raises ``TransformError``; never anything else for bad input."""
    if len(content) > MAX_INPUT_BYTES:
        raise TransformError("The file is larger than %d bytes." % MAX_INPUT_BYTES)
    problems = validate(steps)
    if problems:
        raise TransformError(problems[0])
    if output not in OUTPUTS:
        raise TransformError("Unknown output format %r." % output)
    try:
        records = READERS[steps[0]["type"]]["read"](content, steps[0])
        records_in = len(records)
        ctx = Context(lookups=lookups or {})
        for step in steps[1:]:
            records = STEPS[step["type"]]["apply"](records, step, ctx)
        data = write_output(records, output, root=root, record=record)
    except TransformError:
        raise
    except (KeyError, TypeError, AttributeError, IndexError, RecursionError, MemoryError) as exc:
        raise TransformError("The steps do not fit this file (%s)." % type(exc).__name__) from None
    content_type, extension, _ = OUTPUTS[output]
    return Result(data, content_type, extension, records_in, len(records), ctx.stats, records[:PREVIEW_RECORDS])


# -- index tables (the data behind the "lookup" step) ---------------------------------------------------------------------------------------------------

MAX_LOOKUP_BYTES = 10 * 1000 * 1000
MAX_LOOKUP_ROWS = 100_000
MAX_LOOKUP_COLUMNS = 60


def parse_lookup(content: bytes, filename: str = "", sheet: str | None = None) -> tuple[list[str], list[dict]]:
    columns, rows, _ = parse_lookup_ex(content, filename, sheet)
    return columns, rows


def parse_lookup_ex(content: bytes, filename: str = "", sheet: str | None = None) -> tuple[list[str], list[dict], str]:
    """An uploaded index table as (columns, rows, a note about how it was read): Excel (.xlsx, one sheet), CSV with
    a header row, or a JSON list of objects."""
    if len(content) > MAX_LOOKUP_BYTES:
        raise TransformError("The index file is larger than %d MB." % (MAX_LOOKUP_BYTES // 1_000_000))
    kind = detect_format(content, filename)
    note = ""
    if kind == "xlsx_records":
        names = xlsx_sheets(content)
        used, table = read_xlsx_table(content, sheet or None, max_rows=MAX_LOOKUP_ROWS + 1)
        header, body = _header_and_rows(table)
        columns = header
        rows = [{c: v for c, v in zip(columns, r)} for r in body]
        note = "Read from sheet %r%s." % (used, "; the other %d sheet%s ignored" % (len(names) - 1, " was" if len(names) == 2 else "s were") if len(names) > 1 else "")
    elif kind == "json_records":
        data = _json_load(content)
        if isinstance(data, dict):
            data = next((v for v in data.values() if isinstance(v, list)), None)
        if not isinstance(data, list) or not all(isinstance(r, dict) for r in data):
            raise TransformError("The JSON must be a list of objects, one per row.")
        columns: list[str] = []
        for r in data:
            for k in r:
                if k not in columns:
                    columns.append(str(k))
        rows = [{c: _text(_json_scalar(r.get(c))) for c in columns} for r in data]
    elif kind == "xml_records":
        raise TransformError("Upload the index as Excel (.xlsx), CSV or JSON.")
    else:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            try:
                text = content.decode("cp1253")  # Greek Windows exports
            except UnicodeDecodeError:
                raise TransformError("Cannot read the file: save it as UTF-8.") from None
        sample = text[:4096]
        delimiter = ";" if sample.count(";") > sample.count(",") else ("\t" if sample.count("\t") > sample.count(",") else ",")
        reader = csv.DictReader(io.StringIO(text, newline=""), delimiter=delimiter)
        if not reader.fieldnames:
            raise TransformError("The CSV has no header row.")
        columns = [c.strip() for c in reader.fieldnames if c is not None]
        rows = [{c.strip(): _text(v) for c, v in r.items() if c is not None} for r in reader]
    if not columns or not rows:
        raise TransformError("The index file has no rows.")
    if len(columns) > MAX_LOOKUP_COLUMNS or len(rows) > MAX_LOOKUP_ROWS:
        raise TransformError("At most %d columns and %d rows." % (MAX_LOOKUP_COLUMNS, MAX_LOOKUP_ROWS))
    if len(set(columns)) != len(columns) or any(not c for c in columns):
        raise TransformError("Every column needs its own non-empty name.")
    return columns, rows, note


def lookup_table(columns: list[str], rows: list[dict]) -> dict:
    """The in-memory form the lookup step uses (it builds its key index on first use)."""
    return {"columns": columns, "rows": rows, "index": {}}
